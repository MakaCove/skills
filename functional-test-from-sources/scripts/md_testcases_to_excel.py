#!/usr/bin/env python3
"""Export an explicitly authored six-table Markdown bundle; never infer coverage."""

from __future__ import annotations

import argparse
import html
import re
import sys
import tempfile
from pathlib import Path

from testcase_io import (
    REQUIRED_SHEETS, SHEET_HEADERS, ensure_input_paths, normalize,
    split_markdown_row,
)
from validate_testcases import validate, validate_workbook


def parse_workbook_markdown(text: str) -> dict[str, list[dict[str, str]]]:
    """Require named sections and exact current headers; reject lossy legacy conversion."""
    tables: dict[str, list[dict[str, str]]] = {}
    current = None
    lines = text.splitlines()
    index = 0
    separator = re.compile(r"^\s*\|(?:\s*:?-{3,}:?\s*\|)+\s*$")
    while index < len(lines):
        heading = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", lines[index])
        if heading:
            title = heading.group(1).strip().strip("`")
            current = title if title in SHEET_HEADERS else None
        if (current and index + 1 < len(lines)
                and lines[index].strip().startswith("|")
                and separator.fullmatch(lines[index + 1])):
            if current in tables:
                raise ValueError(f"{current} 章节有多张表，请先人工整理")
            headers = split_markdown_row(lines[index])
            if headers != SHEET_HEADERS[current]:
                raise ValueError(f"{current} 表头与当前契约不一致；旧格式请先补齐并转为标准六表")
            rows = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = split_markdown_row(lines[index])
                if len(cells) != len(headers):
                    raise ValueError(f"{current} 第 {index + 1} 行列数错误，文字中的竖线请转义")
                cells = [html.unescape(re.sub(r"<br\s*/?>", "\n", cell, flags=re.I)) for cell in cells]
                row = dict(zip(headers, cells))
                if any(normalize(value) for value in row.values()):
                    rows.append(row)
                index += 1
            tables[current] = rows
            continue
        index += 1
    if not tables:
        raise ValueError("未找到标准六表章节；仅有旧功能点/链路表不能自动推导需求、场景或覆盖状态")
    return tables


def write_sheet(worksheet, headers: list[str], rows: list[dict[str, str]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, row in enumerate(rows, start=2):
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = normalize(row.get(header, ""))
            # Test input such as '=1+1' is literal data, not an Excel formula.
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        "内容": 90, "操作步骤": 48, "页面预期结果": 48,
        "接口或数据预期结果": 48, "前置条件": 48, "测试数据": 30,
        "备注": 42, "说明": 48, "覆盖说明": 48, "缺少的信息": 48,
        "影响范围": 30, "建议确认内容": 48, "关联需求": 24, "对应用例": 36,
    }
    for column, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = widths.get(header, 18)
    worksheet.row_dimensions[1].height = 28


def write_fixed_workbook(workbook, tables: dict[str, list[dict[str, str]]]) -> None:
    missing = set(REQUIRED_SHEETS) - tables.keys()
    if missing:
        raise ValueError(f"缺少独立整理的工作表: {', '.join(sorted(missing))}；不能从已有用例反推")
    for sheet in REQUIRED_SHEETS:
        write_sheet(workbook.create_sheet(sheet), SHEET_HEADERS[sheet], tables[sheet])


def export_markdown(paths: list[Path], output: Path) -> None:
    import openpyxl

    tables: dict[str, list[dict[str, str]]] = {}
    for path in paths:
        for sheet, rows in parse_workbook_markdown(path.read_text(encoding="utf-8-sig")).items():
            tables.setdefault(sheet, []).extend(rows)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        write_fixed_workbook(workbook, tables)
        errors = validate_workbook(workbook, at_generation=True)
        if errors:
            raise ValueError("导出前校验失败：\n" + "\n".join(errors))
        if output.suffix.lower() != ".xlsx":
            raise ValueError("输出必须为 .xlsx")
        output.parent.mkdir(parents=True, exist_ok=True)
        # Save to an owned temporary file so a failure cannot overwrite an existing deliverable.
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
        try:
            workbook.save(temporary)
            errors = validate(temporary, at_generation=True)
            if errors:
                raise ValueError("导出后校验失败：\n" + "\n".join(errors))
            temporary.replace(output)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="搬运已完整整理的标准六表 Markdown 为 Excel，不自动推导覆盖")
    parser.add_argument("md_files", nargs="+", help="含命名章节与标准表头的内部 Markdown 文件")
    parser.add_argument("-o", "--output", default=None, help="输出 .xlsx 路径")
    args = parser.parse_args()
    try:
        paths = ensure_input_paths(args.md_files, {".md"})
        output = Path(args.output) if args.output else paths[0].with_suffix(".xlsx")
        export_markdown(paths, output)
    except ImportError:
        print("错误: 缺少 openpyxl，导出与校验未完成", file=sys.stderr)
        return 2
    except (OSError, UnicodeError, ValueError) as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 1
    print(f"已导出: {output}；结构与引用校验通过，业务覆盖仍需人工核对")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
