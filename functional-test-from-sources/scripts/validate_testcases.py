#!/usr/bin/env python3
"""Check workbook structure and references, not business correctness or execution."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from testcase_io import (
    CANONICAL_METHODS, CASE_ID, CASE_TYPES, COVERAGE_STATUS, DB_VALUES,
    EXECUTION_RESULTS, GAP_ID, INFO_KEYS, METHOD_STATUS, PRIORITIES,
    REQUIRED_SHEETS, REQUIREMENT_ID, SCENARIO_ID, SHEET_HEADERS,
    find_case_references, normalize, split_methods, split_references,
)


def rows_by_header(ws) -> tuple[list[str], list[dict[str, str]]]:
    values = ws.iter_rows(values_only=True)
    first = next(values, None)
    if first is None:
        return [], []
    headers = [normalize(value) for value in first]
    rows = []
    for raw in values:
        row = {header: normalize(raw[i]) if i < len(raw) else ""
               for i, header in enumerate(headers) if header}
        if any(row.values()):
            rows.append(row)
    return headers, rows


def issue(errors: list[str], code: str, message: str) -> None:
    errors.append(f"ERROR [{code}] {message}")


def references(text: str, pattern, errors: list[str], label: str, required: bool = False) -> set[str]:
    tokens = split_references(text)
    if required and not tokens:
        issue(errors, "EMPTY_REFERENCE", f"{label} 没有编号")
    invalid = [token for token in tokens if not pattern.fullmatch(token)]
    if invalid:
        issue(errors, "INVALID_REFERENCE", f"{label} 含非法编号或范围缩写: {', '.join(invalid)}")
    if len(tokens) != len(set(tokens)):
        issue(errors, "DUPLICATE_REFERENCE", f"{label} 重复列举编号")
    return {token for token in tokens if pattern.fullmatch(token)}


def required_values(row: dict[str, str], fields, errors: list[str], label: str) -> None:
    for field in fields:
        if not row.get(field):
            issue(errors, "EMPTY_REQUIRED", f"{label} 缺少字段: {field}")


def validate_workbook(workbook, at_generation: bool = False) -> list[str]:
    errors: list[str] = []
    if workbook.sheetnames != REQUIRED_SHEETS:
        issue(errors, "SHEET_ORDER", f"工作表必须按 {REQUIRED_SHEETS} 排列，实际为 {workbook.sheetnames}")

    parsed = {}
    for sheet, expected in SHEET_HEADERS.items():
        if sheet not in workbook.sheetnames:
            parsed[sheet] = []
            continue
        headers, rows = rows_by_header(workbook[sheet])
        parsed[sheet] = rows
        if headers != expected:
            issue(errors, "HEADERS", f"{sheet} 表头及顺序必须为 {expected}，实际为 {headers}")

    for sheet in ("测试场景", "详细测试用例", "需求覆盖"):
        if not parsed[sheet]:
            issue(errors, "EMPTY_SHEET", f"{sheet} 没有实际内容；不能以空表宣布完成")

    info = {}
    for row in parsed["测试说明"]:
        key = row.get("项目", "")
        required_values(row, ("项目", "内容"), errors, "测试说明")
        if key in info:
            issue(errors, "DUPLICATE_INFO", f"测试说明项目重复: {key}")
        info[key] = row.get("内容", "")
    for key in INFO_KEYS:
        if not info.get(key):
            issue(errors, "MISSING_INFO", f"测试说明缺少: {key}")

    cases: dict[str, dict[str, str]] = {}
    case_requirements: dict[str, set[str]] = {}
    method_cases = {method: set() for method in CANONICAL_METHODS}
    numbers = []
    for row in parsed["详细测试用例"]:
        case_id = row.get("用例编号", "")
        if not CASE_ID.fullmatch(case_id):
            issue(errors, "INVALID_ID", f"用例编号非法: {case_id or '<空>'}")
        else:
            numbers.append(int(case_id.split("-")[1]))
        if case_id in cases:
            issue(errors, "DUPLICATE_ID", f"用例编号重复: {case_id}")
        cases[case_id] = row
        required_values(row, (
            "模块", "用例标题", "设计方法", "优先级", "类型", "前置条件",
            "测试数据", "操作步骤", "页面预期结果", "执行结果",
        ), errors, case_id or "<空>")
        if row.get("优先级") not in PRIORITIES:
            issue(errors, "INVALID_PRIORITY", f"{case_id} 优先级必须为 P0 至 P3")
        if row.get("类型") not in CASE_TYPES:
            issue(errors, "INVALID_TYPE", f"{case_id} 类型必须为正常、边界或异常")
        methods = split_methods(row.get("设计方法", ""))
        if any(method not in CANONICAL_METHODS for method in methods) or len(methods) != len(set(methods)):
            issue(errors, "INVALID_METHOD", f"{case_id} 方法须使用固定名称，多个用加号且不得重复")
        for method in methods:
            if method in method_cases:
                method_cases[method].add(case_id)
        if row.get("数据库验证") not in DB_VALUES:
            issue(errors, "INVALID_DATABASE_CHECK", f"{case_id} 数据库验证必须为需要或不需要")
        data_expected = row.get("接口或数据预期结果", "")
        if (row.get("页面预期结果", "").startswith("不适用") or row.get("数据库验证") == "需要"):
            if not data_expected or data_expected.startswith("不适用"):
                issue(errors, "MISSING_DATA_EXPECTATION", f"{case_id} 必须填写实际接口或数据断言")
        if row.get("执行结果") not in EXECUTION_RESULTS:
            issue(errors, "INVALID_EXECUTION_STATE", f"{case_id} 执行结果必须为未执行、通过或失败")
        if at_generation and row.get("执行结果") != "未执行":
            issue(errors, "INVALID_GENERATION_STATE", f"{case_id} 生成时执行结果必须为未执行")
        case_requirements[case_id] = references(
            row.get("关联需求", ""), REQUIREMENT_ID, errors, f"{case_id} 关联需求", required=True,
        )

    if numbers and sorted(numbers) != list(range(1, len(numbers) + 1)):
        issue(errors, "NON_CONTIGUOUS_IDS", "用例编号必须从 TC-001 连续递增，不能重复或跳号")
    valid_ids = {case_id for case_id in cases if CASE_ID.fullmatch(case_id)}

    method_rows = parsed["测试方法覆盖"]
    if [row.get("设计方法", "") for row in method_rows] != CANONICAL_METHODS:
        issue(errors, "METHOD_COVERAGE_ROWS", "方法覆盖必须按八种固定方法的顺序各一行")
    for row in method_rows:
        method = row.get("设计方法", "")
        status = row.get("适用性", "")
        refs = references(row.get("对应用例", ""), CASE_ID, errors, f"{method} 对应用例")
        if status not in METHOD_STATUS:
            issue(errors, "INVALID_METHOD_STATUS", f"{method} 适用性非法")
        if status == "适用" and not refs:
            issue(errors, "MISSING_METHOD_CASES", f"{method} 标记适用但没有用例")
        if status in {"不适用", "适用但未覆盖"} and not row.get("说明"):
            issue(errors, "MISSING_METHOD_REASON", f"{method} 必须说明不适用原因或具体缺口")
        for ref in sorted(refs - valid_ids):
            issue(errors, "UNKNOWN_METHOD_CASE", f"{method} 引用不存在的用例: {ref}")
        actual = method_cases.get(method, set())
        if status == "不适用" and (refs or actual):
            issue(errors, "METHOD_STATUS_CONFLICT", f"{method} 不适用却引用或实际用于用例")
        if method in method_cases and refs != actual:
            issue(errors, "METHOD_REFERENCE_MISMATCH",
                  f"{method} 方法覆盖引用与详细用例不一致；应为 {sorted(actual)}，实际为 {sorted(refs)}")

    requirements: dict[str, set[str]] = {}
    for row in parsed["需求覆盖"]:
        req_id = row.get("需求编号", "")
        if not REQUIREMENT_ID.fullmatch(req_id):
            issue(errors, "INVALID_REQUIREMENT_ID", f"需求覆盖每行须为一个 R-n 或 AC-n: {req_id or '<空>'}")
        if req_id in requirements:
            issue(errors, "DUPLICATE_REQUIREMENT", f"需求编号重复: {req_id}")
        required_values(row, ("需求摘要", "说明"), errors, req_id or "<空需求>")
        refs = references(row.get("对应用例", ""), CASE_ID, errors, f"{req_id} 对应用例")
        requirements[req_id] = refs
        status = row.get("覆盖状态", "")
        if status not in COVERAGE_STATUS:
            issue(errors, "INVALID_COVERAGE_STATUS", f"{req_id} 覆盖状态非法")
        if status in {"已覆盖", "部分覆盖"} and not refs:
            issue(errors, "COVERAGE_WITHOUT_CASES", f"{req_id} 标记{status}却没有用例")
        if status == "未覆盖" and refs:
            issue(errors, "UNCOVERED_WITH_CASES", f"{req_id} 已有用例，不能标未覆盖")
        for ref in sorted(refs):
            if ref not in valid_ids:
                issue(errors, "UNKNOWN_COVERAGE_ID", f"{req_id} 引用不存在的用例: {ref}")
            elif req_id not in case_requirements.get(ref, set()):
                issue(errors, "REQUIREMENT_REFERENCE_MISMATCH", f"{req_id} 引用 {ref}，但该用例未关联此需求")
    for case_id, reqs in case_requirements.items():
        for req_id in sorted(reqs):
            if req_id not in requirements:
                issue(errors, "UNKNOWN_REQUIREMENT", f"{case_id} 引用清单中不存在的需求: {req_id}")
            elif case_id not in requirements[req_id]:
                issue(errors, "MISSING_REVERSE_TRACE", f"{case_id} 关联 {req_id}，但需求覆盖未列举该用例")

    seen_scenarios = set()
    assigned_cases = set()
    for row in parsed["测试场景"]:
        scene_id = row.get("场景编号", "")
        if not SCENARIO_ID.fullmatch(scene_id) or scene_id in seen_scenarios:
            issue(errors, "INVALID_SCENARIO_ID", f"场景编号非法或重复: {scene_id or '<空>'}")
        seen_scenarios.add(scene_id)
        required_values(row, ("场景名称", "覆盖说明"), errors, scene_id or "<空场景>")
        if row.get("优先级") not in PRIORITIES:
            issue(errors, "INVALID_SCENARIO_PRIORITY", f"{scene_id} 优先级非法")
        reqs = references(row.get("关联需求", ""), REQUIREMENT_ID, errors, f"{scene_id} 关联需求", required=True)
        for req_id in sorted(reqs - requirements.keys()):
            issue(errors, "UNKNOWN_SCENARIO_REQUIREMENT", f"{scene_id} 引用不存在的需求: {req_id}")
        explanation = row.get("覆盖说明", "")
        refs = find_case_references(explanation)
        if not refs:
            issue(errors, "MISSING_SCENARIO_CASES", f"{scene_id} 覆盖说明必须显式列举用例编号")
        if re.search(r"TC-[0-9]+\s*(?:~|～|至|到|—|–|-)\s*(?:TC-)?[0-9]+", explanation):
            issue(errors, "SCENARIO_RANGE", f"{scene_id} 用例范围须逐个展开")
        covered_reqs = set()
        for ref in sorted(refs):
            if ref not in valid_ids:
                issue(errors, "UNKNOWN_SCENARIO_CASE", f"{scene_id} 引用不存在的用例: {ref}")
                continue
            assigned_cases.add(ref)
            case_reqs = case_requirements.get(ref, set())
            covered_reqs.update(case_reqs)
            if not reqs.intersection(case_reqs):
                issue(errors, "SCENARIO_REQUIREMENT_MISMATCH", f"{scene_id} 与 {ref} 无共同关联需求")
        if reqs - covered_reqs:
            issue(errors, "SCENARIO_UNMAPPED_REQUIREMENT", f"{scene_id} 有需求未对应所列用例: {sorted(reqs - covered_reqs)}")
    for case_id in sorted(valid_ids - assigned_cases):
        issue(errors, "CASE_WITHOUT_SCENARIO", f"{case_id} 未归属任何场景")

    seen_gaps = set()
    for row in parsed["信息缺口"]:
        gap_id = row.get("编号", "")
        if not GAP_ID.fullmatch(gap_id) or gap_id in seen_gaps:
            issue(errors, "INVALID_GAP_ID", f"缺口编号非法或重复: {gap_id or '<空>'}")
        seen_gaps.add(gap_id)
        required_values(row, SHEET_HEADERS["信息缺口"], errors, gap_id or "<空缺口>")
        if not re.match(r"^[是否](?:$|[，,；;：:\s])", row.get("是否影响执行", "")):
            issue(errors, "INVALID_GAP_IMPACT", f"{gap_id} 是否影响执行须以是/否开头，可跟说明")

    return errors


def validate(path: Path, at_generation: bool = False) -> list[str]:
    try:
        import openpyxl
    except ImportError:
        return ["ERROR [DEPENDENCY] 缺少 openpyxl，校验未执行"]
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        return [f"ERROR [INPUT] 无法读取 Excel: {exc}"]
    try:
        return validate_workbook(workbook, at_generation=at_generation)
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="校验六表测试用例的结构与引用，不证明业务覆盖完整或执行通过")
    parser.add_argument("workbook", help="功能测试用例.xlsx")
    parser.add_argument("--generation", action="store_true", help="生成阶段：执行结果必须全部为未执行")
    args = parser.parse_args()
    path = Path(args.workbook)
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        print("ERROR [INPUT] 请输入存在的 .xlsx 文件", file=sys.stderr)
        return 2
    errors = validate(path, at_generation=args.generation)
    for item in errors:
        print(item, file=sys.stderr)
    if errors:
        print(f"校验失败: {len(errors)} 个问题", file=sys.stderr)
        return 1
    print("结构与引用校验通过；仍需人工核对业务断言与覆盖范围，不代表业务测试已执行")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
