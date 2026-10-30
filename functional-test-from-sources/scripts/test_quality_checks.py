#!/usr/bin/env python3
"""Regression checks for delivery contracts and conservative source inventory."""
from __future__ import annotations

import copy
import tempfile
import unittest
from pathlib import Path

import openpyxl

from analyze_codebase import analyze_frontend, correlate, regex_endpoints
from md_testcases_to_excel import (
    export_markdown, parse_workbook_markdown, write_fixed_workbook,
)
from testcase_io import SHEET_HEADERS, escape_markdown_cell
from validate_testcases import validate, validate_workbook

EXAMPLE = Path(__file__).resolve().parents[1] / "references" / "示例.md"


def make_workbook(tables):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    write_fixed_workbook(workbook, tables)
    return workbook


def markdown(tables):
    lines = []
    for sheet, rows in tables.items():
        headers = SHEET_HEADERS[sheet]
        lines.extend([f"## {sheet}", "", "| " + " | ".join(headers) + " |",
                      "|" + "|".join("---" for _ in headers) + "|"])
        for row in rows:
            lines.append("| " + " | ".join(escape_markdown_cell(row.get(h, "")) for h in headers) + " |")
        lines.append("")
    return "\n".join(lines)


class ContractChecks(unittest.TestCase):
    def setUp(self):
        self.tables = copy.deepcopy(parse_workbook_markdown(EXAMPLE.read_text(encoding="utf-8")))

    def errors(self):
        workbook = make_workbook(self.tables)
        try:
            return validate_workbook(workbook, at_generation=True)
        finally:
            workbook.close()

    def assert_error(self, code):
        self.assertTrue(any(f"[{code}]" in item for item in self.errors()), self.errors())

    def test_complete_example_with_unresolved_requirement_is_valid(self):
        self.assertEqual(self.errors(), [])

    def test_empty_cases_scenarios_and_requirements_fail(self):
        for name in ("详细测试用例", "测试场景", "需求覆盖"):
            self.tables[name] = []
        for row in self.tables["测试方法覆盖"]:
            row.update(适用性="不适用", 对应用例="", 说明="本次示例不适用")
        self.assert_error("EMPTY_SHEET")

    def test_unknown_method_case_fails(self):
        self.tables["测试方法覆盖"][0]["对应用例"] = "TC-999"
        self.assert_error("UNKNOWN_METHOD_CASE")

    def test_method_must_match_actual_case_tags(self):
        self.tables["测试方法覆盖"][0]["对应用例"] = "TC-001、TC-002"
        self.assert_error("METHOD_REFERENCE_MISMATCH")

    def test_used_method_cannot_be_not_applicable(self):
        self.tables["测试方法覆盖"][0].update(适用性="不适用", 对应用例="", 说明="错误判断")
        self.assert_error("METHOD_STATUS_CONFLICT")

    def test_empty_or_unknown_method_tokens_fail(self):
        for value in ("+", "等价类 +", "等价类划分", "边界值分析", "等价类 + 等价类"):
            with self.subTest(value=value):
                self.tables["详细测试用例"][0]["设计方法"] = value
                self.assert_error("INVALID_METHOD")

    def test_partial_method_can_reference_existing_cases_and_explain_gap(self):
        self.tables["测试方法覆盖"][0].update(适用性="适用但未覆盖", 说明="尚缺另一个已识别等价类")
        self.assertEqual(self.errors(), [])

    def test_unknown_requirement_fails(self):
        self.tables["详细测试用例"][0]["关联需求"] = "R-999"
        self.assert_error("UNKNOWN_REQUIREMENT")

    def test_missing_reverse_trace_fails(self):
        self.tables["需求覆盖"][0]["对应用例"] = "TC-002"
        self.assert_error("MISSING_REVERSE_TRACE")

    def test_wrong_requirement_to_case_link_fails(self):
        self.tables["需求覆盖"][1]["对应用例"] += "、TC-002"
        self.assert_error("REQUIREMENT_REFERENCE_MISMATCH")

    def test_covered_requirement_requires_cases(self):
        self.tables["需求覆盖"][-1]["覆盖状态"] = "已覆盖"
        self.assert_error("COVERAGE_WITHOUT_CASES")

    def test_explicit_uncovered_requirement_is_preserved(self):
        self.tables["需求覆盖"][-1]["覆盖状态"] = "未覆盖"
        self.assertEqual(self.errors(), [])

    def test_unknown_coverage_case_fails(self):
        self.tables["需求覆盖"][0]["对应用例"] += "、TC-999"
        self.assert_error("UNKNOWN_COVERAGE_ID")

    def test_unknown_scenario_case_fails(self):
        self.tables["测试场景"][0]["覆盖说明"] += "、TC-999"
        self.assert_error("UNKNOWN_SCENARIO_CASE")

    def test_orphan_case_fails(self):
        self.tables["测试场景"][1]["覆盖说明"] = "TC-006、TC-007"
        self.assert_error("CASE_WITHOUT_SCENARIO")

    def test_scenario_range_is_not_silently_accepted(self):
        self.tables["测试场景"][0]["覆盖说明"] = "TC-001~TC-005"
        self.assert_error("SCENARIO_RANGE")

    def test_scenario_requirement_must_exist(self):
        self.tables["测试场景"][0]["关联需求"] += "、R-999"
        self.assert_error("UNKNOWN_SCENARIO_REQUIREMENT")

    def test_db_assertion_requires_data_expectation(self):
        self.tables["详细测试用例"][0].update(数据库验证="需要", 接口或数据预期结果="")
        self.assert_error("MISSING_DATA_EXPECTATION")

    def test_interface_only_case_requires_data_expectation(self):
        self.tables["详细测试用例"][4]["接口或数据预期结果"] = ""
        self.assert_error("MISSING_DATA_EXPECTATION")

    def test_generation_rejects_executed_results_but_later_validation_allows_them(self):
        self.tables["详细测试用例"][0]["执行结果"] = "通过"
        self.assert_error("INVALID_GENERATION_STATE")
        workbook = make_workbook(self.tables)
        try:
            self.assertEqual(validate_workbook(workbook), [])
        finally:
            workbook.close()

    def test_duplicate_and_noncontiguous_ids_fail(self):
        self.tables["详细测试用例"][1]["用例编号"] = "TC-001"
        self.assert_error("DUPLICATE_ID")
        self.assert_error("NON_CONTIGUOUS_IDS")

    def test_missing_source_information_fails(self):
        self.tables["测试说明"] = []
        self.assert_error("MISSING_INFO")

    def test_unknown_method_alias_and_reference_range_fail(self):
        self.tables["测试方法覆盖"][0]["对应用例"] = "TC-001~TC-005"
        self.assert_error("INVALID_REFERENCE")

    def test_wrong_sheet_order_and_duplicate_headers_fail(self):
        workbook = make_workbook(self.tables)
        try:
            workbook._sheets.reverse()
            workbook["详细测试用例"].cell(1, 2, "用例编号")
            errors = validate_workbook(workbook, at_generation=True)
            self.assertTrue(any("[SHEET_ORDER]" in item for item in errors))
            self.assertTrue(any("[HEADERS]" in item for item in errors))
        finally:
            workbook.close()

    def test_real_excel_round_trip_preserves_literal_test_inputs(self):
        self.tables["详细测试用例"][0]["测试数据"] = "=1+1\n甲|乙 & <输入>"
        with tempfile.TemporaryDirectory(prefix="functional-skill-check-") as folder:
            source = Path(folder) / "input.md"
            target = Path(folder) / "output.xlsx"
            source.write_text(markdown(self.tables), encoding="utf-8")
            export_markdown([source], target)
            self.assertEqual(validate(target, at_generation=True), [])
            workbook = openpyxl.load_workbook(target, data_only=False)
            try:
                cell = workbook["详细测试用例"]["I2"]
                self.assertEqual(cell.value, "=1+1\n甲|乙 & <输入>")
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(workbook["需求覆盖"]["A6"].value, "R-5")
            finally:
                workbook.close()

    def test_incomplete_conversion_does_not_overwrite_existing_output(self):
        tables = {"详细测试用例": self.tables["详细测试用例"]}
        with tempfile.TemporaryDirectory(prefix="functional-skill-check-") as folder:
            source = Path(folder) / "input.md"
            target = Path(folder) / "output.xlsx"
            source.write_text(markdown(tables), encoding="utf-8")
            target.write_bytes(b"existing deliverable")
            with self.assertRaises(ValueError):
                export_markdown([source], target)
            self.assertEqual(target.read_bytes(), b"existing deliverable")

    def test_legacy_details_cannot_fabricate_coverage(self):
        with self.assertRaises(ValueError):
            parse_workbook_markdown("| 模块 | ID | 标题 |\n|---|---|---|\n| A | TC-001 | 保存 |")


class InventoryChecks(unittest.TestCase):
    def endpoints(self, source):
        return regex_endpoints(Path("source"), Path("source") / "Controller.java", source)

    def test_spring_class_prefix_is_not_an_independent_endpoint(self):
        endpoints = self.endpoints('''
        @RestController
        @RequestMapping("/order")
        class Orders {
          @PostMapping("/save")
          public Object save() { return null; }
          @GetMapping
          public Object list() { return null; }
        }
        ''')
        self.assertEqual([(e["method"], e["normalized_path"]) for e in endpoints],
                         [("POST", "/order/save"), ("GET", "/order")])

    def test_spring_path_arrays_and_request_method(self):
        endpoints = self.endpoints('''
        @RequestMapping(path={"/a", "/b"}) class A {
          @RequestMapping(value={"/x", "/y"}, method=RequestMethod.POST)
          public Object save() { return null; }
        }
        ''')
        self.assertEqual({e["normalized_path"] for e in endpoints},
                         {"/a/x", "/a/y", "/b/x", "/b/y"})
        self.assertTrue(all(e["method"] == "POST" for e in endpoints))

    def test_multiple_classes_do_not_share_prefixes(self):
        endpoints = self.endpoints('''
        @RequestMapping("/a") class A {
          @GetMapping("/x") public Object x() { return null; }
        }
        @RequestMapping("/b") class B {
          @GetMapping("/y") public Object y() { return null; }
        }
        ''')
        self.assertEqual([e["normalized_path"] for e in endpoints], ["/a/x", "/b/y"])

    def test_spring_constants_are_not_claimed_resolved(self):
        for source in (
            '@RequestMapping(BASE) class A { @GetMapping("/x") Object x() { return null; } }',
            'class A { @GetMapping(PATH) Object x() { return null; } }',
            '@RequestMapping("${base}") class A { @GetMapping("/x") Object x() { return null; } }',
        ):
            with self.subTest(source=source):
                endpoints = self.endpoints(source)
                self.assertTrue(endpoints)
                self.assertTrue(all(e["needs_review"] and e["normalized_path"] is None for e in endpoints))

    def test_comment_annotations_are_not_endpoints(self):
        endpoints = self.endpoints('// @GetMapping("/fake")\nclass A { @GetMapping("/real") Object x(){return null;} }')
        self.assertEqual([e["normalized_path"] for e in endpoints], ["/real"])

    def test_dynamic_frontend_and_wrapper_default_require_review(self):
        with tempfile.TemporaryDirectory(prefix="functional-scan-check-") as folder:
            root = Path(folder)
            (root / "api.js").write_text("""
            request({url:'/items'});
            request.get('/items/' + id);
            request.get(dynamicPath);
            request.post('/items');
            """, encoding="utf-8")
            frontend = analyze_frontend(root)
            self.assertEqual(len(frontend), 4)
            self.assertEqual(sum(bool(r["needs_review"]) for r in frontend), 3)
            backend = [{"normalized_path": "/items", "method": "POST"}]
            matches = correlate(frontend, backend)
            self.assertEqual(sum(m["status"] == "matched_candidate" for m in matches), 1)

    def test_unmatched_scan_is_not_a_missing_function_conclusion(self):
        matches = correlate([{"normalized_path": "/items", "method": "POST"}], [])
        self.assertEqual(matches[0]["status"], "not_matched_in_scan")

    def test_dynamic_fetch_method_is_not_assumed_get(self):
        with tempfile.TemporaryDirectory(prefix="functional-scan-check-") as folder:
            root = Path(folder)
            (root / "api.js").write_text("fetch('/items', {method: chosenMethod});", encoding="utf-8")
            requests = analyze_frontend(root)
            self.assertEqual(requests[0]["method"], "UNKNOWN")
            self.assertTrue(requests[0]["needs_review"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
